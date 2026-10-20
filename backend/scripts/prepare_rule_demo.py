"""Build an independent, reproducible synthetic demo; never overwrite existing data."""
import argparse
import hashlib
import json
from pathlib import Path
import re
import sqlite3
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.config import PROJECT_ROOT
from scripts.init_database import CUSTOMERS, INVOICES


def prepare(target: Path):
    if target.exists():
        raise FileExistsError(f"目标已存在，拒绝覆盖：{target}")
    from docx import Document
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    target.mkdir(parents=True)
    docs = target / "documents"
    docs.mkdir()
    refunds = [(1, 9, 10.0), (2, 9, 5.0), (3, 10, 5.0), (4, 15, 3.0)]
    with sqlite3.connect(target / "business.db") as db:
        db.executescript("""
        PRAGMA foreign_keys=ON;
        CREATE TABLE Customer(CustomerId INTEGER PRIMARY KEY, FirstName TEXT, LastName TEXT, Country TEXT);
        CREATE TABLE Invoice(InvoiceId INTEGER PRIMARY KEY, CustomerId INTEGER NOT NULL REFERENCES Customer,
            InvoiceDate TEXT NOT NULL, BillingCountry TEXT NOT NULL, Total REAL NOT NULL CHECK(Total >= 0));
        CREATE TABLE Refund(RefundId INTEGER PRIMARY KEY, InvoiceId INTEGER NOT NULL REFERENCES Invoice,
            Amount REAL NOT NULL CHECK(Amount >= 0));
        CREATE INDEX idx_invoice_date ON Invoice(InvoiceDate);
        CREATE INDEX idx_refund_invoice ON Refund(InvoiceId);
        """)
        db.executemany("INSERT INTO Customer VALUES (?,?,?,?)", CUSTOMERS)
        db.executemany("INSERT INTO Invoice VALUES (?,?,?,?,?)", INVOICES)
        db.executemany("INSERT INTO Refund VALUES (?,?,?)", refunds)
    (docs / "销售指标口径说明.md").write_bytes((PROJECT_ROOT / "data/documents/销售指标口径说明.md").read_bytes())
    regional = (PROJECT_ROOT / "data/documents/地区经营建议.md").read_text(encoding="utf-8")
    for name, text in re.findall(r"## (\w+)\s*\n(.*?)(?=\n## |\Z)", regional, re.S):
        (docs / f"{name}经营建议.md").write_text(f"# {name}\n\n{text.strip()}", encoding="utf-8")
    for name, text in {
        "客户统计说明.txt": "客户数量按照CustomerId去重统计，同名客户不能合并。",
        "发票时间说明.txt": "年度发票统计以InvoiceDate为时间依据，开始日期包含，结束日期不包含。",
        "数据使用说明.txt": "本数据包为人工构造的合成演示资料，仅用于机制验证，不代表真实业务表现或竞赛得分。",
    }.items():
        (docs / name).write_text(text, encoding="utf-8")
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    pdf = canvas.Canvas(str(docs / "退款计算规则.pdf"), pagesize=(595, 842), invariant=1)
    pdf.setTitle("退款计算规则 - 合成演示")
    pdf.setFillColorRGB(.04, .25, .43)
    pdf.setFont("STSong-Light", 22)
    pdf.drawString(55, 765, "退款计算规则")
    pdf.setFillColorRGB(.15, .15, .15)
    pdf.setFont("STSong-Light", 13)
    pdf.drawString(55, 690, "净收入等于销售额减去已确认退款金额。")
    pdf.setFont("STSong-Light", 11)
    pdf.drawString(55, 620, "适用年度2012。")
    pdf.drawString(55, 70, "合成演示样本 / 第1页")
    pdf.save()
    doc = Document()
    doc.add_heading("客户资格计算条款", 0)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "优惠资格要求同一年度已开具发票销售额至少80元，按CustomerId汇总。"
    doc.save(docs / "优惠资格规则.docx")
    metadata = {name: {"valid_from": "2012-01-01", "valid_to": "2013-01-01", "version": "synthetic-v1"}
                for name in ("退款计算规则.pdf", "优惠资格规则.docx")}
    (docs / "source_metadata.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    manifest = {"synthetic": True, "customer_count": len(CUSTOMERS), "invoice_count": len(INVOICES),
                "confirmed_refunds": refunds,
                "semantics": "Refund.Amount仅为已确认退款；按原始发票年度归属，扣除其全部关联退款",
                "oracle": {"2012_USA_sales": 160, "2012_USA_refund": 20, "2012_USA_net_income": 140,
                           "2012_eligible_customer_ids_at_80": [2, 9]},
                "files": {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in docs.iterdir() if p.is_file()}}
    (target / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def activate(target):
    # Keep unrelated configuration, including credentials, private and unchanged.
    env = PROJECT_ROOT / ".env"
    text = env.read_text(encoding="utf-8") if env.exists() else ""
    values = {"DATABASE_PATH": target / "business.db", "DOCUMENTS_PATH": target / "documents",
              "INDEXES_PATH": target / "indexes", "PARSED_PATH": target / "parsed",
              "DEMO_DATA_LABEL": "合成演示数据：10位客户、20张发票、4笔已确认退款；仅用于机制验证"}
    for key, value in values.items():
        replacement = f'{key}="{str(value).replace(chr(92), chr(47))}"'
        pattern = rf"(?m)^{key}=.*$"
        text = re.sub(pattern, lambda _: replacement, text) if re.search(pattern, text) else text.rstrip() + "\n" + replacement + "\n"
    env.write_text(text, encoding="utf-8")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--target", type=Path, default=PROJECT_ROOT / "data/rule_demo")
    parser.add_argument("--activate", action="store_true")
    parser.add_argument("--activate-existing", action="store_true", help="仅连接已有合成包，不重建或覆盖数据")
    args = parser.parse_args()
    if args.activate_existing:
        target = args.target.resolve()
        manifest = json.loads((target / "manifest.json").read_text(encoding="utf-8"))
        if manifest.get("synthetic") is not True or not (target / "business.db").is_file() or not (target / "documents").is_dir():
            parser.error("不是完整的合成演示包")
        activate(target)
        print("已连接现有合成演示包；请重启后端以刷新配置。")
        raise SystemExit(0)
    manifest = prepare(args.target.resolve())
    if args.activate:
        activate(args.target.resolve())
    print(json.dumps({"target": str(args.target), "documents": len(manifest["files"]) - 1,
                      "oracle": manifest["oracle"], "activated": args.activate}, ensure_ascii=False))
